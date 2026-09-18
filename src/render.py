import sys, os, html, importlib.util
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

HERE=os.path.dirname(os.path.abspath(__file__))
DECK_OUT=os.environ.get("DECK_OUT",os.path.join(HERE,"..","decks"))
HTML_OUT=os.environ.get("TEMPLATE_OUT",os.path.join(HERE,"..","docs","templates"))
spec=importlib.util.spec_from_file_location('decks',HERE+'/decks.py'); decks=importlib.util.module_from_spec(spec); spec.loader.exec_module(decks)

PAPER=RGBColor(0xF7,0xF9,0xFB); WHITE=RGBColor(0xFF,0xFF,0xFF); INK=RGBColor(0x1B,0x1F,0x2A)
INK2=RGBColor(0x4B,0x55,0x63); INK3=RGBColor(0x7B,0x84,0x94); RED=RGBColor(0xC4,0x12,0x30)
REDSOFT=RGBColor(0xF9,0xE3,0xE7); GRID=RGBColor(0xD3,0xDE,0xEA); RULE=RGBColor(0xB8,0xC6,0xD6)
DISPLAY="Georgia"; BODY="Calibri"; MONO="Consolas"

# ======================================================================= PPTX
def build_pptx(deck, out):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    W=prs.slide_width; H=prs.slide_height; BLANK=prs.slide_layouts[6]
    def bg(slide, color=PAPER):
        f=slide.background.fill; f.solid(); f.fore_color.rgb=color
    def grid(slide, step=0.35):
        x=0
        while x<=W:
            ln=slide.shapes.add_connector(1,Emu(x),0,Emu(x),H); ln.line.color.rgb=GRID; ln.line.width=Pt(.5); x+=Inches(step)
        y=0
        while y<=H:
            ln=slide.shapes.add_connector(1,0,Emu(y),W,Emu(y)); ln.line.color.rgb=GRID; ln.line.width=Pt(.5); y+=Inches(step)
    def rect(slide,x,y,w,h,fill=WHITE,line=None,lw=.75):
        s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,x,y,w,h); s.fill.solid(); s.fill.fore_color.rgb=fill
        if line is None: s.line.fill.background()
        else: s.line.color.rgb=line; s.line.width=Pt(lw)
        s.shadow.inherit=False; return s
    def hline(slide,x,y,w,color=RULE,lw=.75):
        ln=slide.shapes.add_connector(1,x,y,x+w,y); ln.line.color.rgb=color; ln.line.width=Pt(lw); return ln
    def text(slide,x,y,w,h,paras,anchor=MSO_ANCHOR.TOP):
        tb=slide.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        first=True
        for p in paras:
            para=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
            para.alignment=p.get('align',PP_ALIGN.LEFT)
            if 'space_after' in p: para.space_after=Pt(p['space_after'])
            if 'line' in p: para.line_spacing=p['line']
            if p.get('bullet'):
                pPr=para._p.get_or_add_pPr(); pPr.set('marL',str(Inches(.22))); pPr.set('indent',str(-Inches(.22)))
                bu=etree.SubElement(pPr,qn('a:buClr')); c=etree.SubElement(bu,qn('a:srgbClr')); c.set('val','C41230')
                bf=etree.SubElement(pPr,qn('a:buFont')); bf.set('typeface','Arial')
                bc=etree.SubElement(pPr,qn('a:buChar')); bc.set('char','•')
            for r in (p.get('runs') or [p]):
                run=para.add_run(); run.text=r['t']; f=run.font
                f.name=r.get('font',p.get('font',BODY)); f.size=Pt(r.get('size',p.get('size',16)))
                f.color.rgb=r.get('color',p.get('color',INK)); f.bold=r.get('bold',p.get('bold',False)); f.italic=r.get('italic',p.get('italic',False))
        return tb
    def label(slide,x,y,w,t,color=INK3,size=11):
        return text(slide,x,y,w,Inches(.3),[{'t':t.upper(),'font':MONO,'size':size,'color':color}])
    def footer(slide,n,section):
        hline(slide,Inches(.6),H-Inches(.55),W-Inches(1.2),RULE)
        text(slide,Inches(.6),H-Inches(.5),Inches(8),Inches(.3),[{'t':f"Miami University · Math Ed Portfolio · {deck['name']}",'font':MONO,'size':9,'color':INK3}])
        text(slide,W-Inches(3.6),H-Inches(.5),Inches(3),Inches(.3),[{'t':f'{section}   {n:02d}'.strip(),'font':MONO,'size':9,'color':INK3,'align':PP_ALIGN.RIGHT}])
    def title_block(slide,eyebrow,title,sub=None,tw=Inches(9.5)):
        label(slide,Inches(.6),Inches(.55),Inches(8),eyebrow,RED)
        text(slide,Inches(.6),Inches(.85),tw,Inches(1),[{'t':title,'font':DISPLAY,'size':30,'color':INK}])
        if sub: text(slide,Inches(.6),Inches(1.6),Inches(11.5),Inches(.8),[{'t':sub,'size':14,'color':INK2,'line':1.1}])
    def bullets(items,hs=14,bs=13):
        return [{'runs':[{'t':a+'  ','size':hs,'bold':True},{'t':b,'size':bs,'color':INK2}],'bullet':True,'space_after':7,'line':1.1} for a,b in items]

    for i,sl in enumerate(deck['slides']):
        n=i+1; t=sl['type']
        s=prs.slides.add_slide(BLANK); bg(s)
        if t=='title':
            grid(s)
            rect(s,Inches(.9),Inches(1.2),Inches(7.6),Inches(5.1),WHITE); rect(s,Inches(.9),Inches(1.2),Inches(.08),Inches(5.1),RED)
            label(s,Inches(1.35),Inches(1.6),Inches(6),sl['eyebrow'],RED)
            text(s,Inches(1.35),Inches(1.95),Inches(6.8),Inches(2.4),[{'runs':[{'t':sl['head']+'\n','font':DISPLAY,'size':42,'color':INK},{'t':sl['head2'],'font':DISPLAY,'size':42,'color':RED,'italic':True}]}])
            text(s,Inches(1.35),Inches(4.35),Inches(6.6),Inches(1.2),[{'t':sl['lede'],'size':15,'color':INK2,'line':1.1}])
            text(s,Inches(1.35),Inches(5.6),Inches(6.6),Inches(.5),[{'runs':[{'t':'Selection. ','font':DISPLAY,'size':18},{'t':'Reflection. ','font':DISPLAY,'size':18,'color':RED,'bold':True},{'t':'Connection.','font':DISPLAY,'size':18}]}])
            rect(s,Inches(9.0),Inches(1.2),Inches(3.5),Inches(5.1),WHITE,RULE)
            text(s,Inches(9.25),Inches(1.45),Inches(3.0),Inches(4.7),[
              {'runs':[{'t':'Definition 1.1 ','font':DISPLAY,'size':13,'bold':True},{'t':'(ePortfolio).','font':DISPLAY,'size':13,'italic':True}],'space_after':4},
              {'t':'A curated collection of student-created work — for learning, assessment, and showcase — with reflection on each artifact and on the connections between them.','font':DISPLAY,'size':12,'color':INK2,'space_after':10},
              {'runs':[{'t':'Corollary ','font':DISPLAY,'size':13,'bold':True},{'t':'(Math Ed Portfolio).','font':DISPLAY,'size':13,'italic':True}],'space_after':4},
              {'t':'The same thing, organized around what a beginning mathematics teacher needs to be able to do. Without reflection it is a repository. With it, it is the argument, with evidence, that you are ready for your own classroom.','font':DISPLAY,'size':12,'color':INK2,'space_after':8},
              {'t':'□','font':DISPLAY,'size':12,'color':INK3,'align':PP_ALIGN.RIGHT}])
            text(s,Inches(.9),H-Inches(.5),Inches(8),Inches(.3),[{'t':deck['name'].upper()+'  ·  '+deck['who'].upper(),'font':MONO,'size':9,'color':INK3}])
            continue
        footer(s,n,sl['eyebrow'].split(' of ')[0])
        title_block(s,sl['eyebrow'],sl['head'],sl.get('sub'))
        if t=='features':
            x0=Inches(.6); y0=Inches(2.55); cw=Inches(2.9); gap=Inches(.2)
            for j,(h,b) in enumerate(sl['items']):
                x=x0+j*(cw+gap); rect(s,x,y0,cw,Inches(3.85),WHITE,RULE); rect(s,x,y0,cw,Inches(.06),RED)
                text(s,x+Inches(.22),y0+Inches(.3),cw-Inches(.44),Inches(.9),[{'t':h,'font':DISPLAY,'size':18}])
                text(s,x+Inches(.22),y0+Inches(1.15),cw-Inches(.44),Inches(2.6),[{'t':b,'size':12.5,'color':INK2,'line':1.15}])
        elif t=='arc':
            ly=Inches(3.35); hline(s,Inches(.6),ly,W-Inches(1.2),RULE,1.5); cw=Inches(2.95); gap=Inches(.12)
            for j,(y,h,b) in enumerate(sl['stages']):
                x=Inches(.6)+j*(cw+gap)
                dot=s.shapes.add_shape(MSO_SHAPE.OVAL,x,ly-Inches(.11),Inches(.22),Inches(.22)); dot.fill.solid(); dot.fill.fore_color.rgb=RED if j==3 else PAPER; dot.line.color.rgb=RED; dot.line.width=Pt(1.5); dot.shadow.inherit=False
                label(s,x,ly-Inches(.55),cw,y,INK3)
                text(s,x,ly+Inches(.3),cw-Inches(.2),Inches(.6),[{'t':h,'font':DISPLAY,'size':17}])
                text(s,x,ly+Inches(.85),cw-Inches(.25),Inches(1.7),[{'t':b,'size':12.5,'color':INK2,'line':1.15}])
            rect(s,Inches(.6),Inches(5.75),W-Inches(1.2),Inches(.85),REDSOFT)
            text(s,Inches(.85),Inches(5.8),W-Inches(1.7),Inches(.75),[{'runs':[{'t':sl['recurring'][0]+'  ','font':MONO,'size':11,'color':RED},{'t':sl['recurring'][1],'size':12.5,'color':INK}]}],anchor=MSO_ANCHOR.MIDDLE)
        elif t=='stage':
            lx=Inches(.6); ly=Inches(2.5); lw=Inches(5.6)
            label(s,lx,ly,lw,"What you add",INK3); hline(s,lx,ly+Inches(.32),lw,RULE)
            text(s,lx,ly+Inches(.45),lw,Inches(3.1),bullets(sl['adds'],13,12))
            rx=Inches(6.6); rw=Inches(6.15)
            rect(s,rx,ly-Inches(.12),rw,Inches(3.5),WHITE,RULE); rect(s,rx,ly-Inches(.12),Inches(.06),Inches(3.5),RED)
            label(s,rx+Inches(.3),ly,rw,"What keeps learning going",RED)
            text(s,rx+Inches(.3),ly+Inches(.42),rw-Inches(.55),Inches(2.95),bullets(sl['learn'],13,12))
            rect(s,Inches(.6),Inches(6.1),W-Inches(1.2),Inches(.7),REDSOFT)
            text(s,Inches(.85),Inches(6.15),W-Inches(1.7),Inches(.6),[{'runs':[{'t':'Reflection prompt  ','font':MONO,'size':11,'color':RED},{'t':sl['prompt'],'font':DISPLAY,'size':13,'italic':True,'color':INK}]}],anchor=MSO_ANCHOR.MIDDLE)
        elif t=='cycle':
            x0=Inches(.6); y0=Inches(2.55); cw=Inches(2.95); gap=Inches(.13)
            for j,(h,b) in enumerate(sl['items']):
                x=x0+j*(cw+gap); hline(s,x,y0,cw,RULE); label(s,x,y0+Inches(.12),cw,h,RED,12)
                text(s,x,y0+Inches(.5),cw-Inches(.2),Inches(2.5),[{'t':b,'size':12.5,'color':INK2,'line':1.15}])
            rect(s,Inches(.6),Inches(5.35),W-Inches(1.2),Inches(1.35),WHITE,RULE)
            text(s,Inches(.9),Inches(5.45),W-Inches(1.8),Inches(1.2),[{'t':sl['note'][0],'font':MONO,'size':11,'color':RED,'space_after':4},{'t':sl['note'][1],'size':12.5,'color':INK2,'line':1.15}])
        elif t=='table':
            cols=sl['cols']; rows=sl['rows']; tx=Inches(.6); ty=Inches(2.55); c0=Inches(2.55); rest=W-Inches(1.2)-c0; cw=int(rest/len(cols))
            avail=(H-Inches(.7))-(ty+Inches(.4)); rh=int(avail/len(rows))
            fs=11.5 if len(rows)<=5 else 10.5
            for j,c in enumerate(cols): label(s,tx+c0+j*cw,ty,cw,c,INK3)
            hline(s,tx,ty+Inches(.32),W-Inches(1.2),RULE,1)
            for i2,(r,cells) in enumerate(rows):
                y=ty+Inches(.4)+i2*rh
                text(s,tx,y,c0-Inches(.15),rh,[{'t':r,'font':DISPLAY,'size':13}],anchor=MSO_ANCHOR.MIDDLE)
                for j in range(len(cols)):
                    text(s,tx+c0+j*cw,y,cw-Inches(.18),rh,[{'t':cells[j],'size':fs,'color':INK2,'line':1.08}],anchor=MSO_ANCHOR.MIDDLE)
                hline(s,tx,y+rh,W-Inches(1.2),RULE,.5)
        elif t=='template':
            px=Inches(.6); py=Inches(2.5); pw=Inches(7.4); ph=Inches(4.2)
            rect(s,px,py,pw,ph,WHITE,RULE)
            text(s,px+Inches(.3),py+Inches(.2),pw-Inches(.6),Inches(.5),[{'runs':[{'t':sl['example'][0],'font':DISPLAY,'size':17},{'t':'    '+sl['example'][1],'font':MONO,'size':10,'color':INK3}]}])
            hline(s,px+Inches(.3),py+Inches(.72),pw-Inches(.6),RULE)
            for j,(h,b) in enumerate(sl['sections']):
                y=py+Inches(.9)+j*Inches(.53)
                text(s,px+Inches(.3),y,Inches(1.7),Inches(.5),[{'t':h.upper(),'font':MONO,'size':10,'color':RED}])
                text(s,px+Inches(2.0),y,pw-Inches(2.3),Inches(.5),[{'t':b,'size':12,'color':INK2}])
            rx=Inches(8.4); rw=Inches(4.35)
            label(s,rx,py,rw,sl['map_title'],INK3); hline(s,rx,py+Inches(.32),rw,RULE)
            paras=[]
            for p,tag in sl['map']:
                runs=[{'t':p,'size':12.5,'color':INK if not p.startswith('   ') else INK2}]
                if tag: runs.append({'t':'   '+tag,'font':MONO,'size':9,'color':RED})
                paras.append({'runs':runs,'space_after':5})
            text(s,rx,py+Inches(.45),rw,Inches(3.7),paras)
        elif t=='todo':
            for j,(h,b) in enumerate(sl['items']):
                y=Inches(2.5)+j*Inches(.78)
                text(s,Inches(.6),y,Inches(.5),Inches(.6),[{'t':str(j+1),'font':DISPLAY,'size':22,'color':RED}])
                text(s,Inches(1.2),y,Inches(4.2),Inches(.7),[{'t':h,'font':DISPLAY,'size':15}])
                text(s,Inches(5.5),y+Inches(.04),Inches(7.2),Inches(.75),[{'t':b,'size':12,'color':INK2,'line':1.1}])
                hline(s,Inches(.6),y+Inches(.7),W-Inches(1.2),RULE,.5)
            text(s,Inches(.6),Inches(6.5),W-Inches(1.2),Inches(.4),[{'t':sl['sources'],'font':MONO,'size':9,'color':INK3}])
    prs.save(out); return len(prs.slides)

# ======================================================================= HTML
CSS = """
  :root{
    --paper:#F7F9FB; --paper-2:#FFFFFF; --grid:#D3DEEA; --ink:#1B1F2A; --ink-2:#4B5563; --ink-3:#7B8494;
    --red:#C41230; --red-ink:#FFFFFF; --red-soft:#F9E3E7; --rule:#B8C6D6;
    --display:"STIX Two Text",Georgia,"Times New Roman",serif; --body:"Source Sans 3","Helvetica Neue",Arial,sans-serif; --mono:"IBM Plex Mono",ui-monospace,Menlo,monospace;
  }
  @media (prefers-color-scheme: dark){ :root:not([data-theme="light"]){
    --paper:#1E2A26; --paper-2:#243430; --grid:#2E3F3A; --ink:#F1EFE6; --ink-2:#C9C7BC; --ink-3:#98A39D; --red:#F06A7E; --red-ink:#1E2A26; --red-soft:#3A2B30; --rule:#4B5F58; } }
  :root[data-theme="dark"]{ --paper:#1E2A26; --paper-2:#243430; --grid:#2E3F3A; --ink:#F1EFE6; --ink-2:#C9C7BC; --ink-3:#98A39D; --red:#F06A7E; --red-ink:#1E2A26; --red-soft:#3A2B30; --rule:#4B5F58; }
  *{box-sizing:border-box}
  body{margin:0; color:var(--ink); font-family:var(--body); font-size:17px; line-height:1.5; background-color:var(--paper);
    background-image:linear-gradient(var(--grid) 1px,transparent 1px),linear-gradient(90deg,var(--grid) 1px,transparent 1px); background-size:28px 28px}
  a{color:inherit} p{margin:0} ul{margin:0; padding-left:20px}
  h1,h2,h3,h4{font-family:var(--display); font-weight:400; line-height:1.12; margin:0; text-wrap:balance}
  .wrap{max-width:1080px; margin:0 auto; padding:0 24px}
  .label{font-family:var(--mono); font-size:12px; letter-spacing:.08em; text-transform:uppercase; color:var(--ink-3)}
  .label.red{color:var(--red)}
  header{padding:18px 0; border-bottom:1px solid var(--rule); background:var(--paper)}
  header .wrap{display:flex; justify-content:space-between; align-items:baseline; gap:16px; flex-wrap:wrap}
  .brand{font-family:var(--display); font-size:19px} .brand b{color:var(--red); font-weight:600}
  header a.back{font-size:15px; color:var(--ink-2); text-decoration:none} header a.back:hover{color:var(--red)}
  .hero{padding:64px 0 48px} .hero .wrap{display:grid; grid-template-columns:minmax(0,1.15fr) minmax(0,.85fr); gap:56px; align-items:start}
  .hero h1{font-size:clamp(36px,5vw,56px)} .hero h1 em{font-style:italic; color:var(--red)}
  .lede{font-size:19px; color:var(--ink-2); margin-top:20px; max-width:34em}
  .who{margin-top:22px; font-family:var(--mono); font-size:12.5px; color:var(--ink-3)}
  .motto{font-family:var(--display); font-size:20px; margin-top:26px} .motto b{color:var(--red); font-weight:600}
  .definition{background:var(--paper-2); border-left:4px solid var(--red); padding:22px 24px; font-family:var(--display); font-size:17px; line-height:1.5; box-shadow:0 1px 0 var(--rule)}
  .definition .dt{font-weight:600} .definition .dt i{font-weight:400} .definition p+p{margin-top:12px}
  .definition .qed{display:block; text-align:right; color:var(--ink-3); font-size:15px; margin-top:8px} .definition .qed::after{content:"\\25A1"}
  section{padding:56px 0; border-top:1px solid var(--rule)} section .wrap{background:var(--paper); padding-top:8px; padding-bottom:8px}
  .sec-head{display:grid; grid-template-columns:150px minmax(0,1fr); gap:24px; margin-bottom:32px; align-items:baseline}
  .sec-head h2{font-size:clamp(26px,3.2vw,36px)} .sec-head p{color:var(--ink-2); margin-top:10px; max-width:40em}
  .sec-head .slideno{display:block; margin-top:6px}
  .cards{display:grid; grid-template-columns:repeat(4,minmax(0,1fr)); gap:16px}
  .card{background:var(--paper-2); border:1px solid var(--rule); border-top:3px solid var(--red); padding:20px 20px 22px}
  .card h3{font-size:20px; margin-bottom:10px} .card p{font-size:15px; color:var(--ink-2)}
  .timeline{display:grid; grid-template-columns:repeat(4,minmax(0,1fr)); gap:0; position:relative}
  .timeline::before{content:""; position:absolute; left:0; right:0; top:11px; height:2px; background:var(--rule)}
  .yr{padding:32px 20px 0 0; position:relative}
  .yr::before{content:""; position:absolute; top:5px; left:0; width:14px; height:14px; border-radius:50%; background:var(--paper); border:2px solid var(--red)}
  .yr.now::before{background:var(--red)} .yr .label{display:block; margin-bottom:6px} .yr h3{font-size:20px; margin-bottom:10px} .yr p{color:var(--ink-2); font-size:15.5px}
  .strip{margin-top:28px; background:var(--red-soft); padding:14px 20px; font-size:15px} .strip b{font-family:var(--mono); font-weight:500; font-size:12px; color:var(--red); margin-right:10px}
  .stage{display:grid; grid-template-columns:minmax(0,1fr) minmax(0,1.1fr); gap:24px; align-items:start}
  .stage .col h4{font-family:var(--mono); font-weight:500; font-size:12px; letter-spacing:.08em; text-transform:uppercase; color:var(--ink-3); padding-bottom:10px; border-bottom:1px solid var(--rule); margin-bottom:14px}
  .stage .col.learn{background:var(--paper-2); border:1px solid var(--rule); border-left:4px solid var(--red); padding:20px 22px} .stage .col.learn h4{color:var(--red)}
  .stage li{font-size:15.5px; color:var(--ink-2); margin-bottom:10px} .stage li b{color:var(--ink); font-weight:600}
  .prompt{margin-top:24px; background:var(--red-soft); padding:14px 20px; font-family:var(--display); font-style:italic; font-size:18px} .prompt b{font-family:var(--mono); font-style:normal; font-weight:500; font-size:12px; color:var(--red); margin-right:10px}
  .cycle{display:grid; grid-template-columns:repeat(4,minmax(0,1fr)); gap:0; border-top:1px solid var(--rule); border-bottom:1px solid var(--rule)}
  .cycle div{padding:20px 20px 20px 0} .cycle div+div{border-left:1px solid var(--rule); padding-left:20px} .cycle p{font-size:15px; color:var(--ink-2); margin-top:8px}
  .note{margin-top:24px; background:var(--paper-2); border:1px solid var(--rule); padding:18px 22px} .note p{color:var(--ink-2); font-size:15.5px; margin-top:6px}
  .tablewrap{overflow-x:auto} table{border-collapse:collapse; width:100%; min-width:720px} th{text-align:left; font-family:var(--mono); font-weight:500; font-size:12px; letter-spacing:.08em; text-transform:uppercase; color:var(--ink-3); padding:0 14px 10px 0; border-bottom:1px solid var(--rule)}
  td{padding:12px 14px 12px 0; border-bottom:1px solid var(--rule); font-size:15px; color:var(--ink-2); vertical-align:top} td:first-child{font-family:var(--display); font-size:17px; color:var(--ink); width:22%}
  .tmpl{display:grid; grid-template-columns:minmax(0,1.6fr) minmax(0,1fr); gap:24px; align-items:start}
  .page{background:var(--paper-2); border:1px solid var(--rule); padding:22px 26px} .page h3{font-size:20px; padding-bottom:12px; border-bottom:1px solid var(--rule); margin-bottom:6px} .page h3 span{font-family:var(--mono); font-size:11px; color:var(--ink-3); margin-left:12px}
  .page dl{display:grid; grid-template-columns:150px minmax(0,1fr); gap:10px 16px; margin:14px 0 0} .page dt{font-family:var(--mono); font-size:11.5px; letter-spacing:.06em; text-transform:uppercase; color:var(--red); padding-top:2px} .page dd{margin:0; font-size:15px; color:var(--ink-2)}
  .map ul{list-style:none; padding:0; margin-top:14px} .map li{padding:6px 0; border-bottom:1px solid var(--rule); font-size:15.5px} .map li.sub{padding-left:18px; color:var(--ink-2)} .map li span{font-family:var(--mono); font-size:11px; color:var(--red); margin-left:10px}
  .todo{border-top:1px solid var(--rule)} .todo div{display:grid; grid-template-columns:40px 300px minmax(0,1fr); gap:16px; padding:16px 0; border-bottom:1px solid var(--rule); align-items:baseline}
  .todo .n{font-family:var(--display); font-size:24px; color:var(--red)} .todo h3{font-size:19px} .todo p{font-size:15px; color:var(--ink-2)}
  .sources{margin-top:22px; font-family:var(--mono); font-size:11.5px; color:var(--ink-3)}
  footer{padding:26px 0 40px; color:var(--ink-3); font-size:14px; background:var(--paper)} footer .wrap{display:flex; justify-content:space-between; gap:16px; flex-wrap:wrap}
  @media (max-width:820px){ .hero .wrap,.stage,.tmpl{grid-template-columns:1fr} .sec-head{grid-template-columns:1fr; gap:8px} .cards,.cycle,.timeline{grid-template-columns:1fr}
    .cycle div+div{border-left:0; border-top:1px solid var(--rule); padding-left:0} .timeline::before{display:none} .yr{padding:16px 0 16px 26px} .yr::before{top:22px}
    .todo div{grid-template-columns:32px minmax(0,1fr)} .todo p{grid-column:2} .page dl{grid-template-columns:1fr} }
"""
E=html.escape
def build_html(deck, out, home_url):
    parts=[]
    parts.append(f"<title>{E(deck['name'])}</title>\n<link rel=\"stylesheet\" href=\"https://fonts.googleapis.com/css2?family=STIX+Two+Text:ital,wght@0,400;0,600;1,400&family=Source+Sans+3:wght@400;600&family=IBM+Plex+Mono:wght@400;500&display=swap\">\n<style>{CSS}</style>")
    parts.append(f'<header><div class="wrap"><div class="brand">Miami University · <b>Math Ed Portfolio</b> · {E(deck["name"])}</div><a class="back" href="{home_url}">← Template gallery</a></div></header>')
    for i,sl in enumerate(deck['slides']):
        n=i+1; t=sl['type']
        if t=='title':
            parts.append(f'''<div class="hero"><div class="wrap"><div>
  <div class="label">{E(sl['eyebrow'])}</div>
  <h1 style="margin-top:14px">{E(sl['head'])}<br><em>{E(sl['head2'])}</em></h1>
  <p class="lede">{E(sl['lede'])}</p>
  <p class="who">Template · {E(deck['who'])}</p>
  <p class="motto">Selection. <b>Reflection.</b> Connection.</p></div>
  <aside class="definition"><p><span class="dt">Definition 1.1 <i>(ePortfolio).</i></span> A curated collection of student-created work — for learning, assessment, and showcase — with reflection on each artifact and on the connections between them.</p>
  <p><span class="dt">Corollary <i>(Math Ed Portfolio).</i></span> The same thing, organized around what a beginning mathematics teacher needs to be able to do. Without reflection it is a repository. With it, it is the argument, with evidence, that you are ready for your own classroom.</p><span class="qed"></span></aside>
</div></div>'''); continue
        head=f'<div class="sec-head"><div><span class="label red">{E(sl["eyebrow"])}</span><span class="label slideno">Slide {n:02d}</span></div><div><h2>{E(sl["head"])}</h2>' + (f'<p>{E(sl["sub"])}</p>' if sl.get('sub') else '') + '</div></div>'
        body=''
        if t=='features':
            body='<div class="cards">'+''.join(f'<div class="card"><h3>{E(h)}</h3><p>{E(b)}</p></div>' for h,b in sl['items'])+'</div>'
        elif t=='arc':
            body='<div class="timeline">'+''.join(f'<div class="yr{" now" if j==3 else ""}"><span class="label">{E(y)}</span><h3>{E(h)}</h3><p>{E(b)}</p></div>' for j,(y,h,b) in enumerate(sl['stages']))+'</div>'
            body+=f'<div class="strip"><b>{E(sl["recurring"][0])}</b>{E(sl["recurring"][1])}</div>'
        elif t=='stage':
            li=lambda items:''.join(f'<li><b>{E(a)}</b> {E(b)}</li>' for a,b in items)
            body=f'<div class="stage"><div class="col"><h4>What you add</h4><ul>{li(sl["adds"])}</ul></div><div class="col learn"><h4>What keeps learning going</h4><ul>{li(sl["learn"])}</ul></div></div>'
            body+=f'<div class="prompt"><b>Reflection prompt</b>{E(sl["prompt"])}</div>'
        elif t=='cycle':
            body='<div class="cycle">'+''.join(f'<div><span class="label red">{E(h)}</span><p>{E(b)}</p></div>' for h,b in sl['items'])+'</div>'
            body+=f'<div class="note"><span class="label red">{E(sl["note"][0])}</span><p>{E(sl["note"][1])}</p></div>'
        elif t=='table':
            body='<div class="tablewrap"><table><thead><tr><th></th>'+''.join(f'<th>{E(c)}</th>' for c in sl['cols'])+'</tr></thead><tbody>'
            body+=''.join('<tr><td>'+E(r)+'</td>'+''.join(f'<td>{E(c)}</td>' for c in cells)+'</tr>' for r,cells in sl['rows'])+'</tbody></table></div>'
        elif t=='template':
            body=f'<div class="tmpl"><div class="page"><h3>{E(sl["example"][0])}<span>{E(sl["example"][1])}</span></h3><dl>'+''.join(f'<dt>{E(h)}</dt><dd>{E(b)}</dd>' for h,b in sl['sections'])+'</dl></div>'
            body+=f'<div class="map"><span class="label">{E(sl["map_title"])}</span><ul>'+''.join(f'<li class="{"sub" if p.startswith("   ") else ""}">{E(p.strip())}'+(f'<span>{E(tag)}</span>' if tag else '')+'</li>' for p,tag in sl['map'])+'</ul></div></div>'
        elif t=='todo':
            body='<div class="todo">'+''.join(f'<div><span class="n">{j+1}</span><h3>{E(h)}</h3><p>{E(b)}</p></div>' for j,(h,b) in enumerate(sl['items']))+'</div>'
            body+=f'<p class="sources">{E(sl["sources"])}</p>'
        parts.append(f'<section id="s{n}"><div class="wrap">{head}{body}</div></section>')
    parts.append(f'<footer><div class="wrap"><span>Mathematics Education Program · Miami University</span><span>PowerPoint version: decks/{E(deck["name"])}.pptx</span></div></footer>')
    open(out,'w').write('\n'.join(parts))

if __name__=='__main__':
    home=sys.argv[1] if len(sys.argv)>1 else '../index.html'
    os.makedirs(DECK_OUT,exist_ok=True); os.makedirs(HTML_OUT,exist_ok=True)
    for d in decks.DECKS:
        n=build_pptx(d, os.path.join(DECK_OUT, f"{d['name']}.pptx"))
        build_html(d, os.path.join(HTML_OUT, f"{d['slug']}.html"), home)
        print(d['slug'], n, 'slides')
